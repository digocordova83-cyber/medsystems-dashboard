from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import json

from openpyxl import load_workbook


INPUT = Path("/home/ubuntu/upload/LeadsRecebidos19.08.xlsx")
SAFE_CATEGORY_FIELDS = [
    "Fonte",
    "Etapa",
    "Pipeline de Vendas",
    "Tecnologia",
    "RD Station",
    "Produto de Interesse",
    "Segmento",
    "UF - PF",
    "Lead repetido",
    "Motivo de declínio - WF0",
]
COVERAGE_FIELDS = [
    "Telefone de trabalho",
    "Informações da fonte",
    "Nome da Empresa",
    "Produto de Interesse",
    "Vendedor Responsável",
]


def text(value):
    return str(value).strip() if value not in (None, "") else "Não informado"


def ui_ranking(entries, limit=9, missing_labels=("Não informado", "-", "undefined")):
    normalized = Counter()
    for item in entries:
        label = item["label"]
        if label in missing_labels:
            label = "Não informado"
        normalized[label] += item["count"]
    ranked = [{"label": label, "count": count} for label, count in normalized.most_common()]
    if len(ranked) <= limit:
        return ranked
    kept = ranked[:limit]
    kept.append({"label": "Outros", "count": sum(item["count"] for item in ranked[limit:])})
    return kept


def scope_for_ui(data):
    categories = {
        field: [{"label": label, "count": count} for label, count in values.most_common()]
        for field, values in data["categories"].items()
    }
    pipeline_rows = categories["Pipeline de Vendas"]
    return {
        "totalLeads": data["records"],
        "pipelines": [
            {"label": row["label"], "brand": "BeautySystems" if row["label"] == "Negócios e Redes" else "Medsystems", "count": row["count"]}
            for row in pipeline_rows
        ],
        "origins": ui_ranking(categories["Fonte"]),
        "stages": ui_ranking(categories["Etapa"]),
        "technologies": ui_ranking(categories["Tecnologia"]),
        "interests": ui_ranking(categories["Produto de Interesse"]),
        "segments": ui_ranking(categories["Segmento"], limit=4),
        "states": ui_ranking(categories["UF - PF"]),
        "declineReasons": ui_ranking([item for item in categories["Motivo de declínio - WF0"] if item["label"] != "Não informado"], limit=10),
        "coverage": [{"label": field, "count": data["coverage"][field]} for field in COVERAGE_FIELDS],
        "dailyByPipeline": [
            {"date": date, "medsystems": values.get("Medsystems", 0), "beautysystems": values.get("Negócios e Redes", 0)}
            for date, values in sorted(data["dailyByPipeline"].items())
        ],
    }


workbook = load_workbook(INPUT, read_only=True, data_only=True)
worksheet = workbook["Base"]
rows = worksheet.iter_rows(values_only=True)
headers = [str(value).strip() if value is not None else "" for value in next(rows)]
indexes = {header: index for index, header in enumerate(headers)}
scope_pipelines = {
    "all": {"Medsystems", "Negócios e Redes"},
    "medsystems": {"Medsystems"},
    "beautysystems": {"Negócios e Redes"},
}
scopes = {
    scope: {
        "categories": {field: Counter() for field in SAFE_CATEGORY_FIELDS},
        "coverage": Counter(),
        "dailyByPipeline": defaultdict(Counter),
        "records": 0,
    }
    for scope in scope_pipelines
}
event_records = 0
total_value = 0.0
records = 0

for row in rows:
    records += 1
    source = text(row[indexes["Fonte"]])
    if source == "Evento":
        event_records += 1
        continue
    created = row[indexes["Criado"]]
    pipeline = text(row[indexes["Pipeline de Vendas"]])
    for scope, pipelines in scope_pipelines.items():
        if pipeline not in pipelines:
            continue
        target = scopes[scope]
        target["records"] += 1
        for field in SAFE_CATEGORY_FIELDS:
            target["categories"][field][text(row[indexes[field]])] += 1
        for field in COVERAGE_FIELDS:
            if row[indexes[field]] not in (None, ""):
                target["coverage"][field] += 1
        if hasattr(created, "strftime"):
            target["dailyByPipeline"][created.strftime("%Y-%m-%d")][pipeline] += 1
        value = row[indexes["Total"]]
        if scope == "all" and isinstance(value, (int, float)):
            total_value += value

result = {
    "file": INPUT.name,
    "recordsBeforeEventExclusion": records,
    "eventRecordsExcluded": event_records,
    "recordsAfterEventExclusion": scopes["all"]["records"],
    "dateRange": {
        "start": min(scopes["all"]["dailyByPipeline"]).replace("-", "/"),
        "end": max(scopes["all"]["dailyByPipeline"]).replace("-", "/"),
    },
    "totalValue": total_value,
    "scopes": {
        scope: {
            "records": data["records"],
            "categories": {
                field: [{"label": label, "count": count} for label, count in values.most_common()]
                for field, values in data["categories"].items()
            },
            "coverage": [{"field": field, "count": data["coverage"][field]} for field in COVERAGE_FIELDS],
            "dailyByPipeline": [
                {"date": date, **dict(values)}
                for date, values in sorted(data["dailyByPipeline"].items())
            ],
        }
        for scope, data in scopes.items()
    },
    "uiScopes": {scope: scope_for_ui(data) for scope, data in scopes.items()},
    "dashboardScopes": {
        scope: {
            key: scope_for_ui(data)[key]
            for key in ["totalLeads", "pipelines", "origins", "stages", "dailyByPipeline"]
        }
        for scope, data in scopes.items()
    },
}
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
