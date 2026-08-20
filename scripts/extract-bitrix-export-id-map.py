import json
from pathlib import Path

from openpyxl import load_workbook


input_path = Path("/home/ubuntu/upload/LeadsRecebidos19.08.xlsx")
output_path = Path("/tmp/bitrix-export-id-map.json")
workbook = load_workbook(input_path, read_only=True, data_only=True)
worksheet = workbook["Base"]
rows = worksheet.iter_rows(values_only=True)
headers = [str(value).strip() if value is not None else "" for value in next(rows)]
id_index = headers.index("ID")
pipeline_index = headers.index("Pipeline de Vendas")

records = {}
for row in rows:
    record_id = row[id_index]
    pipeline = row[pipeline_index]
    if record_id not in (None, "") and pipeline not in (None, ""):
        records[str(record_id).strip()] = str(pipeline).strip()

output_path.write_text(json.dumps(records), encoding="utf-8")
print(json.dumps({"records": len(records), "output": str(output_path)}))
