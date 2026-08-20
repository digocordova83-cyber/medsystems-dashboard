import json
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

base = Path('/home/ubuntu/exports')
data = json.loads((base / 'leads_2026-08-01_a_05.json').read_text())
destination = base / 'amostra_leads_RD_e_Bitrix_2026-08-01_a_05.xlsx'

wb = Workbook()
readme = wb.active
readme.title = 'LEIA-ME'
readme.append(['Amostra de leads — 01 a 05/08/2026'])
readme.append(['Fonte RD Station', 'Um registro por contato com o primeiro evento de conversão encontrado no período. Não é cruzado com Bitrix24.'])
readme.append(['Fonte Bitrix24', 'Todos os registros da entidade Lead criados no período. Não é cruzado com RD Station.'])
readme.append(['Fuso do recorte', 'Brasil (GMT-3): 01/08/2026 00:00 até 06/08/2026 00:00, exclusivo.'])
readme.append(['Preservação', 'Nenhuma deduplicação entre fontes foi aplicada.'])

header_fill = PatternFill('solid', fgColor='0F3B52')
header_font = Font(color='FFFFFF', bold=True)

def value(v):
    if isinstance(v, str):
        return v
    if v is None:
        return None
    return str(v)

def add_sheet(title, rows):
    sheet = wb.create_sheet(title)
    if not rows:
        sheet.append(['Sem registros no período'])
        return
    headers = list(rows[0].keys())
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        sheet.append([value(row.get(key)) for key in headers])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for idx, header in enumerate(headers, 1):
        max_len = max(len(str(header)), *(len(str(sheet.cell(row=r, column=idx).value or '')) for r in range(2, sheet.max_row + 1)))
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 42)

add_sheet('RD Station', data['rdRows'])
add_sheet('Bitrix24 Leads', data['bitrixRows'])

for cell in readme[1]:
    cell.fill = header_fill
    cell.font = header_font
readme.column_dimensions['A'].width = 22
readme.column_dimensions['B'].width = 100
readme.freeze_panes = 'A2'
wb.save(destination)
print(destination)
