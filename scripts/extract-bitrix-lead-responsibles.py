import json
from openpyxl import load_workbook

SOURCE = "/home/ubuntu/upload/LeadsRecebidos19.08.xlsx"

workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook["Base"]
rows = sheet.iter_rows(values_only=True)
headers = [str(value).strip() if value is not None else "" for value in next(rows)]
id_index = headers.index("ID")
responsible_index = headers.index("Responsável")

mapping = {}
for row in rows:
    bitrix_id = row[id_index]
    responsible = row[responsible_index]
    if bitrix_id is None or responsible is None:
        continue
    mapping[str(int(bitrix_id))] = str(responsible).strip()

print(json.dumps(mapping, ensure_ascii=False))
