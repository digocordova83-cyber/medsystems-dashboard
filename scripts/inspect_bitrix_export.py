from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


INPUT = Path("/home/ubuntu/upload/LeadsRecebidos19.08.xlsx")
OUTPUT = Path("/home/ubuntu/medsystems-rdstation-oauth/docs/inspecao_exportacao_bitrix_19-08.txt")

wb = load_workbook(INPUT, read_only=True, data_only=True)
lines: list[str] = [f"Arquivo: {INPUT.name}", f"Abas: {', '.join(wb.sheetnames)}", ""]

for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, ())
    headers = [str(value).strip() if value is not None else f"coluna_{index + 1}" for index, value in enumerate(headers)]
    non_empty = Counter()
    categorical = {index: Counter() for index, header in enumerate(headers) if any(token in header.lower() for token in ("pipeline", "funil", "etapa", "status", "origem", "marca", "empresa", "responsável"))}
    count = 0

    for row in rows:
        count += 1
        for index, value in enumerate(row):
            if value not in (None, ""):
                non_empty[index] += 1
                if index in categorical:
                    categorical[index][str(value).strip()] += 1

    lines.extend([f"Aba: {ws.title}", f"Registros: {count}", "Campos:"])
    lines.extend(f"- {header}: {non_empty[index]} preenchidos" for index, header in enumerate(headers))
    if categorical:
        lines.append("Valores categóricos observados:")
        for index, values in categorical.items():
            display = "; ".join(f"{label} ({total})" for label, total in values.most_common(25)) or "sem valores"
            lines.append(f"- {headers[index]}: {display}")
    lines.append("")

OUTPUT.write_text("\n".join(lines), encoding="utf-8")
print(OUTPUT)
